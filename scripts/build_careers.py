"""Build privacy-safe career evidence from genuine OECD PIAAC Cycle 2 respondents.

Only published CSV public-use files are read. No synthetic data, individual data
exports, inferred career histories or country pooling. Results are descriptive.
"""
import csv
import hashlib
import io
import json
import math
import re
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / ".cache" / "piaac"
OUT = ROOT / "data" / "careers.json"
CODEBOOK = ("https://www.oecd.org/content/dam/oecd/en/about/programmes/edu/"
            "piaac/data-materials/cycle-2/piaac-cy2-international-codebook.xlsx")
BASE = "https://webfs.oecd.org/piaac/cy2-puf-data/CSV/"
COUNTRIES = {
    "Chile": "prgchlp2.csv",
    "Poland": "prgpolp2.csv",
    "Singapore": "prgsgpp2.csv",
    "United States": "prgusap2.csv",
    "Italy": "prgitap2.csv",
}
# ISCO-08 TWO-digit occupation families: deliberately not specific job titles.
FAMILIES = {
    "12": "Administrative and commercial managers",
    "13": "Production and specialized services managers",
    "21": "Science and engineering professionals",
    "22": "Health professionals",
    "23": "Teaching professionals",
    "24": "Business and administration professionals",
    "25": "Information and communications technology professionals",
    "26": "Legal, social and cultural professionals",
    "31": "Science and engineering associate professionals",
    "32": "Health associate professionals",
    "33": "Business and administration associate professionals",
    "35": "Information and communications technicians",
}
ED_LEVELS = {"1": "Low", "2": "Medium", "3": "High"}
MIN_RESPONDENTS = 25
MIN_CATEGORY_RESPONDENTS = 20


def download(url, filename):
    CACHE.mkdir(parents=True, exist_ok=True)
    location = CACHE / filename
    if location.exists() and location.stat().st_size > 5000:
        return location
    for attempt in range(4):
        try:
            req = Request(url, headers={"User-Agent": "WhatNextAI/1.0 (OECD PIAAC research demo)"})
            with urlopen(req, timeout=90) as response, (CACHE / (filename + ".partial")).open("wb") as dest:
                while True:
                    data = response.read(1 << 20)
                    if not data:
                        break
                    dest.write(data)
            partial = CACHE / (filename + ".partial")
            if partial.stat().st_size < 5000:
                raise ValueError("Downloaded file is empty or too small")
            partial.replace(location)
            return location
        except (OSError, ValueError) as exc:
            (CACHE / (filename + ".partial")).unlink(missing_ok=True)
            print(f"Download attempt {attempt + 1}: {filename}: {exc}", flush=True)
            if attempt < 3:
                time.sleep(attempt + 1)
    raise RuntimeError(f"Could not download official OECD file: {url}")


def file_sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_codebook():
    from openpyxl import load_workbook
    path = download(CODEBOOK, "piaac-cycle2-codebook.xlsx")
    wb = load_workbook(path, read_only=True, data_only=True)
    entries = {}
    for row in wb["PUF"].iter_rows(values_only=True):
        if len(row) > 8 and row[1] in ("EDFIELD", "ISCO2C", "EDLEVEL3"):
            entries[row[1]] = {
                "description": str(row[2]),
                "codes": str(row[8] or ""),
            }
    if "EDFIELD" not in entries or "ISCO2C" not in entries:
        raise RuntimeError("Cannot validate survey codebook variables.")
    # Codebook has semicolon-separated integer: label entries.
    labels = dict((key, val.strip()) for key, val in re.findall(
        r"(\d{1,2}):\s*([^;]+)", entries["EDFIELD"]["codes"]
    ))
    if len(labels) < 10:
        raise RuntimeError("Unable to parse the official EDFIELD education labels.")
    return labels, {"url": CODEBOOK, "sha256": file_sha256(path)}


def valid_number(value):
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def category_distribution(items):
    """A transparent within-country, within-occupation weighted percentage."""
    weight_sum = sum(items.values())
    if weight_sum <= 0:
        return []
    return [
        {"code": key, "pct": round(weight / weight_sum * 100, 1)}
        for key, weight in sorted(items.items(), key=lambda item: item[1], reverse=True)
    ]


def process_country(name, filename, education_labels):
    url = BASE + filename
    path = download(url, filename)
    summary = defaultdict(lambda: {
        "n": 0, "levels": defaultdict(float), "fields": defaultdict(float),
        "level_n": 0, "field_n": 0,
    })
    with path.open(encoding="utf-8-sig", newline="") as f:
        csv_reader = csv.reader(f, delimiter=";")
        header = next(csv_reader)
        required = ("ISCO2C", "EDFIELD", "EDLEVEL3", "SPFWT0", "C2_D05")
        missing = set(required) - set(header)
        if missing:
            raise RuntimeError(f"{name} PUF missing required fields: {sorted(missing)}")
        ix = {name: header.index(name) for name in required}
        observations = 0
        for values in csv_reader:
            if len(values) <= max(ix.values()):
                continue
            if values[ix["C2_D05"]] != "1":
                continue  # only employed respondents
            family = values[ix["ISCO2C"]].strip()
            if family not in FAMILIES:
                continue  # never invent a category for suppressed ISCO codes
            weight = valid_number(values[ix["SPFWT0"]])
            if weight is None or weight <= 0:
                continue
            item = summary[family]
            item["n"] += 1
            observations += 1
            level = values[ix["EDLEVEL3"]].strip()
            field = values[ix["EDFIELD"]].strip()
            if level in ED_LEVELS:
                item["levels"][level] += weight
                item["level_n"] += 1
            if field in education_labels:
                item["fields"][field] += weight
                item["field_n"] += 1

    groups = {}
    for family, info in sorted(summary.items()):
        if info["n"] < MIN_RESPONDENTS:
            continue
        fields = category_distribution(info["fields"]) if info["field_n"] >= MIN_CATEGORY_RESPONDENTS else []
        levels = category_distribution(info["levels"]) if info["level_n"] >= MIN_CATEGORY_RESPONDENTS else []
        # Don't publish sparse or completely unavailable indicators.
        if not fields and not levels:
            continue
        groups[family] = {
            "title": FAMILIES[family],
            "n": info["n"],
            "education_level_n": info["level_n"],
            "study_field_n": info["field_n"],
            "levels": [
                {"label": ED_LEVELS.get(x["code"], x["code"]), "pct": x["pct"]}
                for x in levels
            ],
            "fields": [
                {"label": education_labels[x["code"]], "pct": x["pct"]}
                for x in fields
            ],
        }
    return {
        "groups": groups,
        "source": url,
        "sha256": file_sha256(path),
        "employed_valid_family_records": observations,
        "published_groups": len(groups),
    }


def main():
    field_labels, codebook = read_codebook()
    results = {}
    excluded = {}
    for country, file in COUNTRIES.items():
        try:
            result = process_country(country, file, field_labels)
            if result["published_groups"] == 0:
                excluded[country] = "No reportable occupation-family groups in this public-use file."
                continue
            results[country] = result
            print(country, result["employed_valid_family_records"],
                  "employed observations /", result["published_groups"], "groups")
        except Exception as exc:
            excluded[country] = type(exc).__name__ + ": " + str(exc)[:160]
            print("Country unavailable:", country, excluded[country], flush=True)
    if len(results) < 2:
        raise RuntimeError("Fewer than two official countries produced valid data; refusing to fabricate results.")
    payload = {
        "dataset": "OECD PIAAC 2023 Cycle 2 Public Use Files",
        "source_url": "https://www.oecd.org/en/data/datasets/piaac-2nd-cycle-database.html",
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "countries": results,
        "excluded": excluded,
        "codebook": codebook,
        "method": "Within each country and ISCO-08 2-digit occupational family: weighted percentages using SPFWT0 on nonmissing education outcomes. No pooled cross-country comparisons. Not replicate-weighted variance estimates.",
        "limitations": [
            "Not individual career histories, proof of causal career success, or exact job qualifications.",
            "Survey respondents are adults in participating countries, not a worldwide random sample.",
            "Education fields and occupations are suppressed or missing in some country files.",
            "Percentages are descriptive and do not include complex-survey uncertainty intervals.",
            "Different categories have different valid-response denominators.",
        ],
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("Saved", OUT, "countries:", len(results), "groups:", sum(len(v["groups"]) for v in results.values()))


if __name__ == "__main__":
    main()
