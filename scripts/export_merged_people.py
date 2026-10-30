"""Export one de-identified merged CSV of *actual* PIAAC observations used.

Five OECD Cycle 2 public-use country files; only employed respondents with valid
ISCO-08 family codes included. No respondent identifiers are retained. Output is
a one-off downloadable GitHub Actions artifact, NOT a public repository commit.
"""
import csv
from pathlib import Path
from scripts.build_careers import (
    COUNTRIES, FAMILIES, ED_LEVELS, download, read_codebook, valid_number
)

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "out" / "WhatNext_Merged_Career_Dataset.csv"

def main():
    field_labels, _ = read_codebook()
    # Verify the exact public-use variables before exporting expanded skill features.
    from openpyxl import load_workbook
    from scripts.build_careers import CODEBOOK, CACHE
    book = load_workbook(CACHE / "piaac-cycle2-codebook.xlsx", read_only=True, data_only=True)
    import re
    for cells in book["PUF"].iter_rows(values_only=True):
        code = str(cells[1] or "") if len(cells) > 1 else ""
        desc = str(cells[2] or "") if len(cells) > 2 else ""
        if ("work" in code.lower() and False) or re.search(r"index of use of|work experience|years.*(working|work|experience)|years.*current|tenure|literacy proficiency|numeracy proficiency|problem solving|learning at work|planning at work|readiness to learn|work autonomy", desc, re.I):
            if any(phrase in desc.lower() for phrase in ["index of use of", "work experience", "years", "tenure", "problem solving", "planning at work", "readiness to learn", "autonomy"]):
                print("CODEBOOK_SKILL",code,desc[:140],flush=True)
    sample_path=download("https://webfs.oecd.org/piaac/cy2-puf-data/CSV/prgchlp2.csv", "prgchlp2.csv")
    with sample_path.open(encoding="utf-8-sig",newline="") as inp:
        fields=next(csv.reader(inp,delimiter=";"))
        print("WORK_VARIABLES", [c for c in fields if re.search(r"WORKC2|JOBEXP|YRS|TENURE|ICTWORK|READWORK|WRITWORK|NUMWORK|TASK|WLE_CA|PVLIT|PVNUM|PVAPS|WORKEX|SKILLUSE|LEARN", c, re.I)][:180],flush=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    headers = [
        "country", "occupation_isco2", "occupation_family",
        "highest_education_field_code", "highest_education_field",
        "education_level_code", "education_level",
        "employment_status", "final_sample_weight",
        "ict_use_at_work_index", "official_source_url"
    ]
    rows_written = 0
    counts = {}
    with OUTPUT.open("w", encoding="utf-8-sig", newline="") as out:
        writer = csv.DictWriter(out, fieldnames=headers)
        writer.writeheader()
        for country, filename in sorted(COUNTRIES.items()):
            original_url = "https://webfs.oecd.org/piaac/cy2-puf-data/CSV/" + filename
            source = download(original_url, filename)
            country_count = 0
            with source.open(encoding="utf-8-sig", newline="") as raw:
                reader = csv.DictReader(raw, delimiter=";")
                required = ("C2_D05", "ISCO2C", "EDFIELD", "EDLEVEL3", "SPFWT0")
                if not set(required).issubset(reader.fieldnames or []):
                    raise RuntimeError(f"Official {country} CSV is missing required columns.")
                for row in reader:
                    if row.get("C2_D05", "").strip() != "1":
                        continue
                    family = row.get("ISCO2C", "").strip()
                    if family not in FAMILIES:
                        continue
                    weight = valid_number(row.get("SPFWT0"))
                    if weight is None or weight <= 0:
                        continue
                    field_code = row.get("EDFIELD", "").strip()
                    level_code = row.get("EDLEVEL3", "").strip()
                    ict = valid_number(row.get("ICTWORKC2"))
                    writer.writerow({
                        "country": country,
                        "occupation_isco2": family,
                        "occupation_family": FAMILIES[family],
                        "highest_education_field_code": field_code if field_code in field_labels else "",
                        "highest_education_field": field_labels.get(field_code, ""),
                        "education_level_code": level_code if level_code in ED_LEVELS else "",
                        "education_level": ED_LEVELS.get(level_code, ""),
                        "employment_status": "Employed",
                        "final_sample_weight": format(weight, ".8g"),
                        "ict_use_at_work_index": format(ict, ".6g") if ict is not None else "",
                        "official_source_url": original_url,
                    })
                    country_count += 1
                    rows_written += 1
            counts[country] = country_count
            print(f"{country}: {country_count} verified respondent records", flush=True)
    if len([n for n in counts.values() if n > 0]) < 2 or rows_written < 100:
        OUTPUT.unlink(missing_ok=True)
        raise RuntimeError("Too few real PIAAC observations: refusing to publish an empty dataset.")
    with OUTPUT.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        assert set(reader.fieldnames or []) == set(headers)
        assert sum(1 for _ in reader) == rows_written
    print(f"MERGED_CSV_READY rows={rows_written} countries={len(counts)} file={OUTPUT}", flush=True)

if __name__ == "__main__":
    main()
