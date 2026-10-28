"""Temporary provenance probe: inspect only public column names and official codebook metadata."""
import csv,io,re,urllib.request
from openpyxl import load_workbook

url="https://www.oecd.org/content/dam/oecd/en/about/programmes/edu/piaac/data-materials/cycle-2/piaac-cy2-international-codebook.xlsx"
try:
    req=urllib.request.Request(url,headers={"User-Agent":"WhatNextAI/1.0 research inspection"})
    with urllib.request.urlopen(req,timeout=45) as f:
        blob=f.read()
    print("CODEBOOK_BYTES",len(blob),flush=True)
    wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    for sh in wb:
        print("SHEET",sh.title,sh.max_row,sh.max_column,flush=True)
        found=0
        for row in sh.iter_rows(values_only=True):
            bits=[str(x)[:150] if x is not None else "" for x in row]
            joined=" | ".join(bits)
            if re.search(r"(ISCED|ISCO|occupational|occupation|education level|field of study|sample weight|employment status|use of.*skill)",joined,re.I):
                print("MATCH",joined[:300],flush=True)
                found+=1
                if found>=45:break
except Exception as e:print("CODEBOOK_ERROR",str(e),flush=True)

for country,file in [("Estonia","prgestp2.csv"),("United States","prgusap2.csv"),("Japan","prgjpnp2.csv")]:
    url="https://webfs.oecd.org/piaac/cy2-puf-data/CSV/"+file
    try:
        req=urllib.request.Request(url,headers={"User-Agent":"WhatNextAI/1.0 research inspection"})
        with urllib.request.urlopen(req,timeout=35) as resp:
            header=resp.readline().decode("utf-8-sig").strip()
            cols=next(csv.reader([header]))
            print("COUNTRY",country,"FIELDS",len(cols),"SIZE",resp.headers.get("Content-Length"),flush=True)
            print("CANDIDATES",[(x) for x in cols if re.search(r"(ISCO|ISCED|EDCAT|FOS|OCC|JOB|EDUC|EMP|W_FNL|SPFWT|READ|NUM|DIG|ICT|PLF|SKILL|V1|PVLIT)",x,re.I)][:120],flush=True)
            print("FIRST_COLUMNS",cols[:75],flush=True)
    except Exception as e:print("COUNTRY_ERROR",country,str(e),flush=True)
